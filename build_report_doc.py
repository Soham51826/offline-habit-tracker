import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_callout_box(doc, text, title="NOTE"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.5)
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Left border styling
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="10B981"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run_title = p.add_run(f"📌 {title}: ")
    run_title.bold = True
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(10)
    run_title.font.color.rgb = RGBColor(16, 185, 129)
    
    run_text = p.add_run(text)
    run_text.font.name = "Calibri"
    run_text.font.size = Pt(10)
    run_text.font.color.rgb = RGBColor(51, 65, 85)

def build_docx(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        md_text = f.read()

    doc = Document()
    
    # Page setup (Letter / A4 with 0.8" margins)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)
    normal_style.paragraph_format.line_spacing = 1.2
    normal_style.paragraph_format.space_after = Pt(6)

    # 1. Title / Cover Page
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(40)
    title_p.paragraph_format.space_after = Pt(10)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = title_p.add_run("A MINI PROJECT REPORT ON\n")
    r_sub.font.size = Pt(13)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(100, 116, 139)

    r_title = title_p.add_run("HabitForge: An Intelligent Hybrid Habit Formation Mobile Application with Gemini AI Behavioral Insights & Google Maps Geofencing\n")
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.name = 'Georgia'
    r_title.font.color.rgb = RGBColor(15, 23, 42)

    p_deg = doc.add_paragraph()
    p_deg.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_deg.paragraph_format.space_before = Pt(16)
    p_deg.paragraph_format.space_after = Pt(24)
    r_d = p_deg.add_run("Submitted in partial fulfillment of the requirements for the award of the Degree of\nBACHELOR OF TECHNOLOGY IN COMPUTER ENGINEERING / INFORMATION TECHNOLOGY\n\nCourse: Mobile Application Development (IT402 / CS402)\nAcademic Year: 2026–2027")
    r_d.font.size = Pt(11)
    r_d.font.italic = True
    r_d.font.color.rgb = RGBColor(71, 85, 105)

    # Candidate Team Table
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    col_widths = [Inches(1.2), Inches(1.8), Inches(3.5)]
    headers = ["Roll No.", "Candidate Name", "Contribution / Focus Area"]
    
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].width = col_widths[i]
        set_cell_background(hdr_cells[i], "1E293B")
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(h)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(10)

    team_data = [
        ("C096", "Soham", "Mobile Architecture, Gemini AI Integration & Database Engine"),
        ("C106", "Student 2", "Google Maps SDK, Geofencing Services & Proximity Triggers"),
        ("C108", "Student 3", "UI/UX Refactoring (Material Design 3), Form Validation & Testing")
    ]
    for roll, name, contrib in team_data:
        row_cells = table.add_row().cells
        for i, val in enumerate([roll, name, contrib]):
            row_cells[i].width = col_widths[i]
            set_cell_background(row_cells[i], "F8FAFC" if i % 2 == 0 else "FFFFFF")
            set_cell_margins(row_cells[i], top=90, bottom=90, left=120, right=120)
            p = row_cells[i].paragraphs[0]
            run = p.add_run(val)
            run.font.size = Pt(9.5)
            if i < 2:
                run.bold = True

    p_dept = doc.add_paragraph()
    p_dept.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_dept.paragraph_format.space_before = Pt(40)
    r_dept = p_dept.add_run("Department of Computer Engineering & Information Technology\nFaculty of Technology & Engineering\nOctober 2026")
    r_dept.font.bold = True
    r_dept.font.size = Pt(11)
    r_dept.font.color.rgb = RGBColor(30, 41, 59)

    doc.add_page_break()

    # Parse remaining markdown sections
    lines = md_text.split('\n')
    idx = 0
    in_table = False
    table_rows = []
    in_code = False
    code_lines = []

    # Skip front page markdown lines (up to Table of Contents)
    start_parsing = False
    for i, line in enumerate(lines):
        if "TABLE OF CONTENTS" in line or "CHAPTER 1" in line:
            idx = i
            start_parsing = True
            break
            
    if not start_parsing:
        idx = 0

    while idx < len(lines):
        line = lines[idx]
        stripped = line.strip()

        # Handle page breaks
        if "\\pagebreak" in stripped:
            doc.add_page_break()
            idx += 1
            continue

        # Code block handling
        if stripped.startswith("```"):
            if in_code:
                # end code block
                code_text = "\n".join(code_lines)
                tbl = doc.add_table(rows=1, cols=1)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                tbl.autofit = False
                tbl.columns[0].width = Inches(6.5)
                c = tbl.cell(0, 0)
                set_cell_background(c, "0F172A")
                set_cell_margins(c, top=100, bottom=100, left=140, right=140)
                cp = c.paragraphs[0]
                crun = cp.add_run(code_text)
                crun.font.name = "Courier New"
                crun.font.size = Pt(8.5)
                crun.font.color.rgb = RGBColor(226, 232, 240)
                in_code = False
                code_lines = []
            else:
                in_code = True
                code_lines = []
            idx += 1
            continue

        if in_code:
            code_lines.append(line)
            idx += 1
            continue

        # Markdown tables
        if stripped.startswith("|") and stripped.endswith("|"):
            if not in_table:
                in_table = True
                table_rows = [stripped]
            else:
                table_rows.append(stripped)
            idx += 1
            continue
        elif in_table:
            # End of table, render it
            in_table = False
            # parse table rows
            parsed_rows = []
            for tr in table_rows:
                if re.match(r'^\|[\s:-]+\|$', tr):
                    continue # separator line
                cols = [c.strip() for c in tr.split("|")[1:-1]]
                parsed_rows.append(cols)

            if parsed_rows:
                num_cols = len(parsed_rows[0])
                t = doc.add_table(rows=len(parsed_rows), cols=num_cols)
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                t.autofit = True
                
                for r_idx, row in enumerate(parsed_rows):
                    for c_idx, cell_val in enumerate(row):
                        if c_idx < len(t.rows[r_idx].cells):
                            c = t.rows[r_idx].cells[c_idx]
                            set_cell_margins(c, top=80, bottom=80, left=100, right=100)
                            p = c.paragraphs[0]
                            p.paragraph_format.space_before = Pt(2)
                            p.paragraph_format.space_after = Pt(2)
                            
                            # Clean markdown bold/tags
                            clean_text = cell_val.replace("**", "").replace("`", "").replace("<br>", "\n")
                            run = p.add_run(clean_text)
                            run.font.size = Pt(9)
                            
                            if r_idx == 0:
                                set_cell_background(c, "1E293B")
                                run.bold = True
                                run.font.color.rgb = RGBColor(255, 255, 255)
                            else:
                                if "PASS" in clean_text:
                                    set_cell_background(c, "ECFDF5")
                                    run.bold = True
                                    run.font.color.rgb = RGBColor(5, 150, 105)
                                else:
                                    set_cell_background(c, "F8FAFC" if r_idx % 2 == 1 else "FFFFFF")
            table_rows = []

        if not stripped:
            idx += 1
            continue

        # Headings
        if stripped.startswith("# "):
            h_text = stripped[2:].strip().replace("**", "")
            h = doc.add_heading(level=1)
            h.paragraph_format.space_before = Pt(20)
            h.paragraph_format.space_after = Pt(8)
            run = h.add_run(h_text)
            run.font.name = "Georgia"
            run.font.size = Pt(17)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42)
        elif stripped.startswith("## "):
            h_text = stripped[3:].strip().replace("**", "")
            h = doc.add_heading(level=2)
            h.paragraph_format.space_before = Pt(16)
            h.paragraph_format.space_after = Pt(6)
            run = h.add_run(h_text)
            run.font.name = "Georgia"
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(30, 41, 59)
        elif stripped.startswith("### "):
            h_text = stripped[4:].strip().replace("**", "")
            h = doc.add_heading(level=3)
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(4)
            run = h.add_run(h_text)
            run.font.name = "Calibri"
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(16, 185, 129)
        elif stripped.startswith("- ") or stripped.startswith("* "):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            clean_line = stripped[2:].strip()
            # parse basic bolding
            parts = re.split(r'(\*\*.*?\*\*)', clean_line)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                else:
                    p.add_run(part)
        elif re.match(r'^\d+\.\s+', stripped):
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(3)
            clean_line = re.sub(r'^\d+\.\s+', '', stripped)
            parts = re.split(r'(\*\*.*?\*\*)', clean_line)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                else:
                    p.add_run(part)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            parts = re.split(r'(\*\*.*?\*\*)', stripped)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                else:
                    p.add_run(part)

        idx += 1

    doc.save(docx_path)
    print(f"[OK] Generated Word Document: {docx_path}")

if __name__ == "__main__":
    md_file = "MINI_PROJECT_REPORT.md"
    docx_file = "MINI_PROJECT_REPORT.docx"
    build_docx(md_file, docx_file)
